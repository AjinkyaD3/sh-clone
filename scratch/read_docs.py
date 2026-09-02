import os
import docx
import openpyxl

task_dir = r"c:\Users\AJINKYA\OneDrive\Desktop\SH NEXT JS\secure-house-nextjs\task"
out_file = r"c:\Users\AJINKYA\OneDrive\Desktop\SH NEXT JS\secure-house-nextjs\scratch\all_docs_dump.txt"

with open(out_file, "w", encoding="utf-8") as out:
    for fname in os.listdir(task_dir):
        fpath = os.path.join(task_dir, fname)
        out.write(f"\n\n=================================\n")
        out.write(f"FILE: {fname}\n")
        out.write(f"=================================\n\n")
        
        if fname.endswith(".docx"):
            try:
                doc = docx.Document(fpath)
                for para in doc.paragraphs:
                    out.write(para.text + "\n")
                out.write("\n[TABLES]\n")
                for table in doc.tables:
                    for row in table.rows:
                        row_data = [cell.text.replace("\n", " ") for cell in row.cells]
                        out.write(" | ".join(row_data) + "\n")
                    out.write("\n")
            except Exception as e:
                out.write(f"Error reading docx: {e}\n")
                
        elif fname.endswith(".xlsx"):
            try:
                wb = openpyxl.load_workbook(fpath, data_only=True)
                for sheet_name in wb.sheetnames:
                    out.write(f"\n--- SHEET: {sheet_name} ---\n")
                    sheet = wb[sheet_name]
                    for row in sheet.iter_rows(values_only=True):
                        row_data = [str(cell).replace('\n', ' ') if cell is not None else "" for cell in row]
                        # Only write rows that have some data
                        if any(row_data):
                            out.write(" | ".join(row_data) + "\n")
            except Exception as e:
                out.write(f"Error reading xlsx: {e}\n")
                
print("Dump completed.")
